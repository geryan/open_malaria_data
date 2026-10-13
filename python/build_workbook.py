"""Build the pilot workbook from the pilot CSVs.

python3 -I python/build_workbook.py <outdir_with_csvs> <extract.json> <rules.tsv> <policy_history.json> <xlsx_out>
"""
import csv, json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

outdir, extract_f, rules_f, hist_f, xlsx_out = sys.argv[1:6]
A = list(csv.DictReader(open(f"{outdir}/articles.csv", encoding="utf-8")))
AU = list(csv.DictReader(open(f"{outdir}/authors.csv", encoding="utf-8")))
D = json.load(open(extract_f, encoding="utf-8"))
RULES = list(csv.DictReader(open(rules_f, encoding="utf-8"), delimiter="\t", quoting=csv.QUOTE_NONE))
H = json.load(open(hist_f, encoding="utf-8"))

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=14, bold=True)
f_link = Font(name=F, size=10, color="0563C1", underline="single")
fill_head = PatternFill("solid", fgColor="305496")
fill_human_head = PatternFill("solid", fgColor="BF8F00")
fill_human = PatternFill("solid", fgColor="FFF2CC")
fill_sub = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")

def as_val(v):
    if v == "True": return True
    if v == "False": return False
    return v

def write_table(ws, headers, rows, widths=None, wrap_cols=(), human_cols=()):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = f_head; c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = fill_human_head if h in human_cols else fill_head
    for i, r in enumerate(rows, 2):
        for j, h in enumerate(headers, 1):
            v = r.get(h, "")
            c = ws.cell(row=i, column=j, value=v)
            c.font = f_norm
            c.alignment = wrap if h in wrap_cols else top
            if h in human_cols: c.fill = fill_human
    ws.freeze_panes = "C2" if len(headers) > 6 else "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for j, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(h, 14)
    ws.row_dimensions[1].height = 42

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active; ws.title = "README"
run = D["run"]
lines = [
    ("Malaria Journal data availability pilot: volume 24 (2025)", f_title),
    ("", None),
    ("What this is", f_bold),
    ("Every item in Malaria Journal volume 24 (articles numbered 1-457, published online 5 Jan to 29 Dec 2025), with title, authors, DOI, corresponding author emails, the verbatim data availability statement (DAS), and a first-pass machine coding of each statement.", None),
    ("The machine coding is the INITIAL coder only. Every row is to be checked by a human in the yellow columns of the 'articles' sheet.", None),
    ("", None),
    ("Sheets", f_bold),
    ("articles: one row per item (457). Machine codes, flags, both statement sources, and human-check columns (yellow).", None),
    ("authors: one row per author (4,533) with affiliations, email where published, corresponding-author flag and ORCID (from Crossref, matched by position).", None),
    ("summary: counts and proportions, calculated by formula from the 'articles' sheet; updates as human codes are entered.", None),
    ("coding_scheme: definitions of categories, flags and derived fields, plus every pattern used (rules/das_coding_rules.tsv).", None),
    ("policy: the journal's current data availability requirements (Oct 2026) and what could be established about their history.", None),
    ("sources_checked: every source consulted, when, how, and the outcome, including what failed.", None),
    ("vol25_published_2025: 68 items published online Nov-Dec 2025 that carry volume 25, excluded by the volume-24 definition.", None),
    ("data_dictionary: every column in 'articles' and 'authors'.", None),
    ("", None),
    ("How to check the coding (yellow columns in 'articles')", f_bold),
    ("human_categories: the correct categories, separated by '; ' in this order: repository; in_article; on_request; restricted_third_party; not_available; not_applicable. Use 'no_statement' or 'unclassified' alone where they apply. Example: on_request; restricted_third_party", None),
    ("human_most_open_route: pick from the drop-down. Example: on_request", None),
    ("human_coder: initials. Example: GR   |   human_date: date checked. Example: 2026-10-20   |   human_notes: free text.", None),
    ("human_agrees_machine: calculated; 'yes' when human_categories exactly matches machine_categories.", None),
    ("", None),
    ("How the data were collected", f_bold),
    (f"Retrieval ran {run['started']} to {run['finished']} (UTC) in the Claude desktop app's browser pane on the user's Mac, because no R runtime or network was available to scripts. Sources: Crossref REST API (metadata, volume), Europe PMC REST API (PMCID, JATS full text), publisher article pages on link.springer.com (article type, authors, affiliations, emails, statement).", None),
    ("The statement is taken from the publisher page (version of record). The Europe PMC JATS statement is kept alongside as a check: after ignoring case, punctuation and doi.org prefixes, the two agree for all 442 items where both exist.", None),
    ("Three statements sat under misspelt headings ('Availiability of data materials', 'Available of data and materials', 'Data Availibility') and were missed by the first heading pattern; they were re-fetched with a revised pattern and verified by SHA-256 against the browser copy (das_patched = TRUE).", None),
    ("Coding: python/build_pilot.py applied rules/das_coding_rules.tsv. R/05_code_das.R implements the same engine on the same rules file for future runs. The R pipeline (R/00-07) has NOT been executed: no R runtime was available. Run R/run_all.R in Positron, then R/07_compare_pilot.R to compare against this pilot.", None),
    ("", None),
    ("Findings needing attention", f_bold),
    ("145 of 457 items (105 of 381 Research articles) give the identical statement 'No datasets were generated or analysed during the current study.' For primary research this is implausible on its face; flag f_na_data_article marks these for checking.", None),
    ("The journal's template for that situation is longer ('Data sharing is not applicable to this article as no datasets were generated or analysed during the current study.'); only 1 item uses it (a Perspective, which also repeats the short sentence). Whether the short sentence is a submission-system default has not been established.", None),
    ("12 items have no statement: 7 corrections, 1 publisher correction, 1 retraction note, 1 editorial, 1 meeting abstract supplement, and 1 review (10.1186/s12936-025-05423-5, a systematic review and meta-analysis).", None),
]
import math
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f or f_norm; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = 20 if f is f_title else 13.5 * max(1, math.ceil(len(t) / 125))
ws.column_dimensions["A"].width = 120

# ---------------------------------------------------------------- articles
CATS = ["repository", "in_article", "on_request", "restricted_third_party", "not_available", "not_applicable"]
FLAGS = [r["code"] for r in RULES if r["kind"] == "flag"]
FLAGS = list(dict.fromkeys(FLAGS)) + ["f_contradiction", "f_na_data_article"]
art_cols = (["article_number", "doi", "published_online", "article_type", "title", "n_authors", "authors",
             "corresponding_authors", "corresponding_emails", "all_author_emails", "das_heading", "das_text",
             "machine_categories", "machine_most_open_route"] + CATS + ["no_statement", "unclassified", "n_categories"] + FLAGS +
            ["das_links", "das_source", "das_patched", "das_sources_agree", "jats_das_text", "epmc_assoc_das_text",
             "code_heading", "code_text", "jats_article_type", "jats_corresponding", "pmcid", "pmid", "url", "update_to",
             "retrieved_utc"])
human_cols = ["human_categories", "human_most_open_route", "human_coder", "human_date", "human_notes", "human_agrees_machine"]
rows = []
for r in A:
    rr = {k: as_val(v) for k, v in r.items()}
    rr["machine_categories"] = r["categories"]
    rr["machine_most_open_route"] = r["most_open_route"]
    rr["article_number"] = int(r["article_number"])
    rr["n_authors"] = int(r["n_authors"])
    rr["n_categories"] = int(r["n_categories"])
    rows.append(rr)
ws = wb.create_sheet("articles")
headers = art_cols + human_cols
widths = {"doi": 30, "title": 50, "authors": 40, "corresponding_authors": 24, "corresponding_emails": 30,
          "all_author_emails": 30, "das_heading": 22, "das_text": 70, "machine_categories": 30,
          "machine_most_open_route": 20, "jats_das_text": 50, "epmc_assoc_das_text": 30, "das_links": 30,
          "code_text": 30, "human_categories": 30, "human_most_open_route": 22, "human_notes": 40,
          "url": 34, "article_type": 14, "published_online": 12, "article_number": 9}
write_table(ws, headers, rows, widths, wrap_cols={"title", "das_text", "jats_das_text", "authors"}, human_cols=set(human_cols))
col = {h: get_column_letter(j) for j, h in enumerate(headers, 1)}
n = len(rows)
for i in range(2, n + 2):
    ws[f"{col['human_agrees_machine']}{i}"] = (f'=IF({col["human_categories"]}{i}="","",'
                                              f'IF(TRIM({col["human_categories"]}{i})={col["machine_categories"]}{i},"yes","no"))')
    ws[f"{col['human_agrees_machine']}{i}"].font = f_norm
    ws[f"{col['human_agrees_machine']}{i}"].fill = fill_human
    ws[f"{col['url']}{i}"].hyperlink = ws[f"{col['url']}{i}"].value
    ws[f"{col['url']}{i}"].font = f_link
    ws.row_dimensions[i].height = 45
routes = ["repository", "in_article", "restricted_third_party", "on_request", "not_available", "not_applicable", "no_statement", "unclassified"]
dv = DataValidation(type="list", formula1='"' + ",".join(routes) + '"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"{col['human_most_open_route']}2:{col['human_most_open_route']}{n + 1}")
ws[f"{col['machine_categories']}1"].comment = Comment("Machine (initial) coding from rules/das_coding_rules.tsv. Multi-label, '; '-separated.", "pilot")
ws[f"{col['machine_most_open_route']}1"].comment = Comment("Most open route claimed, in the order repository > in_article > restricted_third_party > on_request > not_available > not_applicable.", "pilot")

# ---------------------------------------------------------------- authors
ws = wb.create_sheet("authors")
au_rows = [{**r, "position": int(r["position"]), "corresponding": as_val(r["corresponding"])} for r in AU]
write_table(ws, ["doi", "position", "author", "corresponding", "email", "institutions", "orcid"], au_rows,
            {"doi": 30, "author": 28, "email": 32, "institutions": 90, "orcid": 38}, wrap_cols=set())
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- summary
ws = wb.create_sheet("summary", 1)
rng = lambda h: f"articles!${col[h]}$2:${col[h]}${n + 1}"
ws["A1"] = "Summary (formulas on the 'articles' sheet)"; ws["A1"].font = f_title
ws["A3"] = "Items in volume 24"; ws["A3"].font = f_bold
ws["B3"] = f"=COUNTA({rng('doi')})"
r0 = 5
ws.cell(row=r0, column=1, value="Most open route (machine)").font = f_bold
for j, h in enumerate(["route", "n", "% of items", "human n", "% (human)"], 1):
    c = ws.cell(row=r0 + 1, column=j, value=h); c.font = f_head; c.fill = fill_head
for k, rt in enumerate(routes):
    i = r0 + 2 + k
    ws.cell(row=i, column=1, value=rt)
    ws.cell(row=i, column=2, value=f'=COUNTIF({rng("machine_most_open_route")},A{i})')
    ws.cell(row=i, column=3, value=f"=IF($B$3=0,0,B{i}/$B$3)").number_format = "0.0%"
    ws.cell(row=i, column=4, value=f'=COUNTIF({rng("human_most_open_route")},A{i})')
    ws.cell(row=i, column=5, value=f'=IF(SUM($D${r0+2}:$D${r0+1+len(routes)})=0,"",D{i}/SUM($D${r0+2}:$D${r0+1+len(routes)}))').number_format = "0.0%"
tot = r0 + 2 + len(routes)
ws.cell(row=tot, column=1, value="Total").font = f_bold
ws.cell(row=tot, column=2, value=f"=SUM(B{r0+2}:B{tot-1})").font = f_bold
ws.cell(row=tot, column=4, value=f"=SUM(D{r0+2}:D{tot-1})").font = f_bold

r1 = tot + 2
ws.cell(row=r1, column=1, value="Categories claimed (multi-label: an item can have several)").font = f_bold
for j, h in enumerate(["category", "n", "% of items"], 1):
    c = ws.cell(row=r1 + 1, column=j, value=h); c.font = f_head; c.fill = fill_head
for k, ct in enumerate(CATS + ["no_statement", "unclassified"]):
    i = r1 + 2 + k
    ws.cell(row=i, column=1, value=ct)
    ws.cell(row=i, column=2, value=f"=COUNTIF({rng(ct)},TRUE)")
    ws.cell(row=i, column=3, value=f"=IF($B$3=0,0,B{i}/$B$3)").number_format = "0.0%"

r2 = r1 + 2 + len(CATS) + 3
ws.cell(row=r2, column=1, value="Flags").font = f_bold
for j, h in enumerate(["flag", "n", "% of items"], 1):
    c = ws.cell(row=r2 + 1, column=j, value=h); c.font = f_head; c.fill = fill_head
for k, fl in enumerate(FLAGS):
    i = r2 + 2 + k
    ws.cell(row=i, column=1, value=fl)
    ws.cell(row=i, column=2, value=f"=COUNTIF({rng(fl)},TRUE)")
    ws.cell(row=i, column=3, value=f"=IF($B$3=0,0,B{i}/$B$3)").number_format = "0.0%"

r3 = r2 + 2 + len(FLAGS) + 1
types = sorted({r["article_type"] or "(none)" for r in A}, key=lambda t: -sum(1 for r in A if (r["article_type"] or "(none)") == t))
ws.cell(row=r3, column=1, value="Most open route (machine) by article type").font = f_bold
hdr = ["article type", "n"] + routes
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=r3 + 1, column=j, value=h); c.font = f_head; c.fill = fill_head; c.alignment = Alignment(wrap_text=True)
for k, t in enumerate(types):
    i = r3 + 2 + k
    ws.cell(row=i, column=1, value=t)
    ws.cell(row=i, column=2, value=f'=COUNTIF({rng("article_type")},A{i})')
    for m, rt in enumerate(routes):
        ws.cell(row=i, column=3 + m, value=f'=COUNTIFS({rng("article_type")},$A{i},{rng("machine_most_open_route")},{get_column_letter(3+m)}${r3+1})')

r4 = r3 + 2 + len(types) + 1
ws.cell(row=r4, column=1, value="Human check progress").font = f_bold
ws.cell(row=r4 + 1, column=1, value="Rows checked (human_categories filled)")
ws.cell(row=r4 + 1, column=2, value=f'=COUNTIF({rng("human_categories")},"?*")')
ws.cell(row=r4 + 2, column=1, value="Human agrees with machine")
ws.cell(row=r4 + 2, column=2, value=f'=COUNTIF({rng("human_agrees_machine")},"yes")')
ws.cell(row=r4 + 3, column=1, value="Agreement among checked rows")
ws.cell(row=r4 + 3, column=2, value=f'=IF(B{r4+1}=0,"",B{r4+2}/B{r4+1})').number_format = "0.0%"
for row in ws.iter_rows():
    for c in row:
        if c.font.name != F:
            c.font = Font(name=F, size=10, bold=c.font.b)
ws.column_dimensions["A"].width = 44
for j in range(2, 12): ws.column_dimensions[get_column_letter(j)].width = 14

# ---------------------------------------------------------------- coding_scheme
ws = wb.create_sheet("coding_scheme")
defs = [
    ("repository", "category", "Data (or code) deposited in, or said to be openly available from, a named repository, public database or link; accession numbers and data-repository DOIs count. Includes public third-party sources (e.g. DHS) when the statement calls them public."),
    ("in_article", "category", "Data stated to be in the article or its supplementary/additional files."),
    ("on_request", "category", "Data available on request from the authors, corresponding author, PI or study team, or on request with no third party named."),
    ("restricted_third_party", "category", "Access controlled by a third party or process: data access committee, ministry/programme custodian, permission, licence, data transfer/use agreement, registration."),
    ("not_available", "category", "Some or all data explicitly not available, with no route given in the same sentence (often partial, e.g. raw transcripts)."),
    ("not_applicable", "category", "States no data were generated or analysed, or 'not applicable'."),
    ("no_statement", "derived", "No data availability section found on the publisher page or in JATS."),
    ("unclassified", "derived", "A statement exists but no category rule fired; needs human coding."),
    ("most_open_route", "derived", "First category present in the order repository > in_article > restricted_third_party > on_request > not_available > not_applicable."),
    ("f_contradiction", "derived", "not_applicable together with any other category."),
    ("f_na_data_article", "derived", "not_applicable in a Research, Methodology, Case Report or Case Study item."),
]
write_table(ws, ["code", "kind", "definition"], [dict(zip(["code", "kind", "definition"], d)) for d in defs],
            {"code": 26, "kind": 10, "definition": 120}, wrap_cols={"definition"})
start = len(defs) + 4
ws.cell(row=start - 1, column=1, value="Rules (rules/das_coding_rules.tsv). Case-insensitive PCRE. scope = statement or sentence; unless = must not match; require = must also match.").font = f_bold
rh = ["code", "kind", "scope", "pattern", "unless", "require", "description"]
for j, h in enumerate(rh, 1):
    c = ws.cell(row=start, column=j, value=h); c.font = f_head; c.fill = fill_head
for i, r in enumerate(RULES, start + 1):
    for j, h in enumerate(rh, 1):
        c = ws.cell(row=i, column=j, value=r[h]); c.font = Font(name="Courier New", size=9) if h in ("pattern", "unless", "require") else f_norm
        c.alignment = wrap
ws.column_dimensions["D"].width = 90; ws.column_dimensions["E"].width = 45; ws.column_dimensions["F"].width = 30; ws.column_dimensions["G"].width = 45
ws.freeze_panes = "A2"; ws.auto_filter.ref = None

# ---------------------------------------------------------------- policy
ws = wb.create_sheet("policy")
pol = H["policy_rows"]
write_table(ws, ["topic", "summary", "short_quote", "source", "section", "retrieved_utc"], pol,
            {"topic": 26, "summary": 90, "short_quote": 45, "source": 50, "section": 26, "retrieved_utc": 22},
            wrap_cols={"summary", "short_quote", "source"})
ws.freeze_panes = "A2"
hs = len(pol) + 4
ws.cell(row=hs - 1, column=1, value="Archived versions checked (Internet Archive Wayback Machine): presence of key terms in each yearly snapshot").font = f_bold
hh = ["page", "snapshot", "status", "Availability of data and materials", "Availability of supporting data", "All manuscripts must include ...", "'reasonable request' template", "'Data sharing is not applicable' template", "'No datasets were generated or analysed'", "snapshot_url"]
for j, h in enumerate(hh, 1):
    c = ws.cell(row=hs, column=j, value=h); c.font = f_head; c.fill = fill_head; c.alignment = Alignment(wrap_text=True)
for i, s in enumerate(H["snapshots"], hs + 1):
    fl = s.get("flags") or {}
    vals = [s["target"], s["ts"], s.get("status"), fl.get("aodm"), fl.get("aosd"), fl.get("all_manuscripts"), fl.get("reasonable"), fl.get("not_applicable_tmpl"), fl.get("no_datasets"), f"https://web.archive.org/web/{s['ts']}/{s['orig']}"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = f_norm
ws.row_dimensions[hs].height = 60

# ---------------------------------------------------------------- sources_checked
ws = wb.create_sheet("sources_checked")
src_rows = H["sources_checked"]
write_table(ws, ["source", "url_or_query", "what_was_checked", "when_utc", "method", "outcome"], src_rows,
            {"source": 30, "url_or_query": 60, "what_was_checked": 45, "when_utc": 22, "method": 28, "outcome": 70},
            wrap_cols={"url_or_query", "what_was_checked", "outcome", "source"})
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- vol25 list
ws = wb.create_sheet("vol25_published_2025")
v25 = sorted(D["vol25_published_2025"], key=lambda x: int(x["art"]))
v25_rows = [{"doi": x["doi"].lower(), "published_online": "-".join([x["online"].split("-")[0]] + [p.zfill(2) for p in x["online"].split("-")[1:]]), "volume": "25", "article_number": int(x["art"])} for x in v25]
write_table(ws, ["doi", "published_online", "volume", "article_number"], v25_rows, {"doi": 32, "published_online": 16})
ws.freeze_panes = "A2"

# ---------------------------------------------------------------- data dictionary
ws = wb.create_sheet("data_dictionary")
dd = H["data_dictionary"]
write_table(ws, ["sheet", "column", "description", "source"], dd, {"sheet": 10, "column": 26, "description": 100, "source": 34}, wrap_cols={"description"})
ws.freeze_panes = "A2"

wb.save(xlsx_out)
print("saved", xlsx_out)
